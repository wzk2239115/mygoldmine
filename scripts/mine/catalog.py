#!/usr/bin/env python3
import argparse
import concurrent.futures
import datetime as dt
import json
import os
import re
import shutil
import subprocess
import sys
import unicodedata
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_XLSX = ROOT / "agent_llm_audit_catalog.xlsx"
DEFAULT_MANIFEST = ROOT / "catalog" / "manifest.json"
EXISTING_PATHS = {
    "AUD-002": "targets/codex",
    "AUD-061": "targets/sglang",
    "AUD-062": "targets/vllm",
}
URL_RE = re.compile(r"https://github\.com/[A-Za-z0-9_.-]+/[A-Za-z0-9_.-]+")


def now():
    return dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def slugify(audit_id, name):
    ascii_name = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode()
    slug = re.sub(r"[^a-z0-9]+", "-", ascii_name.lower()).strip("-")
    if not slug:
        slug = "project"
    return f"{audit_id.lower()}-{slug}"[:96]


def git_commit(path):
    try:
        return subprocess.run(
            ["git", "-C", str(path), "rev-parse", "HEAD"],
            capture_output=True,
            text=True,
            timeout=20,
            check=True,
        ).stdout.strip()
    except Exception:
        return ""


def load_rows(xlsx):
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        raise SystemExit(f"openpyxl is required: {exc}")
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb["项目清单"]
    rows = list(ws.iter_rows(values_only=True))
    header_index = None
    for index, row in enumerate(rows):
        if row and "编号" in [str(cell).strip() if cell is not None else "" for cell in row]:
            header_index = index
            break
    if header_index is None:
        raise SystemExit("cannot find header row")
    headers = [str(cell).strip() if cell is not None else "" for cell in rows[header_index]]
    result = []
    for row in rows[header_index + 1:]:
        if not row:
            continue
        item = dict(zip(headers, row))
        audit_id = str(item.get("编号") or "").strip()
        if not re.match(r"^AUD-\d+", audit_id):
            continue
        url_match = URL_RE.search(str(item.get("官方仓库与来源") or ""))
        result.append(
            {
                "audit_id": audit_id,
                "batch": str(item.get("建议批次") or "").strip(),
                "direction": str(item.get("方向") or "").strip(),
                "name": str(item.get("项目") or "").replace("\n", " ").strip(),
                "stack": str(item.get("主要技术栈") or "").replace("\n", " ").strip(),
                "entry": str(item.get("建议审计入口") or "").replace("\n", " ").strip(),
                "license": str(item.get("许可证备注") or "").replace("\n", " ").strip(),
                "url": url_match.group(0) if url_match else "",
            }
        )
    wb.close()
    return result


def import_catalog(args):
    rows = load_rows(Path(args.xlsx))
    if args.batch != "all":
        rows = [row for row in rows if row["batch"] == args.batch]
    projects = {}
    for row in rows:
        if not row["url"]:
            continue
        slug = slugify(row["audit_id"], row["name"])
        path = EXISTING_PATHS.get(row["audit_id"], f"targets/catalog/{slug}")
        absolute = ROOT / path
        commit = git_commit(absolute) if absolute.exists() else ""
        projects[slug] = {
            **row,
            "path": path,
            "commit": commit,
            "status": "ok" if commit else "pending",
            "error": "",
        }
    manifest = {
        "schema_version": 1,
        "generated_at": now(),
        "source_catalog": str(Path(args.xlsx).resolve().relative_to(ROOT)),
        "projects": projects,
    }
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    batch_file = out.parent / "first-batch.txt"
    batch_file.write_text(
        "\n".join(
            name for name, item in projects.items() if item["batch"] == "首批候选"
        )
        + "\n"
    )
    print(json.dumps({"manifest": str(out), "projects": len(projects), "first_batch": sum(1 for item in projects.values() if item["batch"] == "首批候选")}, ensure_ascii=False))


def read_manifest(path):
    return json.loads(Path(path).read_text())


def select_projects(manifest, batch):
    items = manifest.get("projects", {})
    if batch == "all":
        return items
    if batch in {"首批候选", "扩展候选", "存量参考"}:
        return {name: item for name, item in items.items() if item.get("batch") == batch}
    wanted = {part.strip() for part in batch.split(",") if part.strip()}
    return {
        name: item
        for name, item in items.items()
        if name in wanted or item.get("audit_id") in wanted
    }


def clone_one(name, item, timeout):
    path = ROOT / item["path"]
    if path.exists() and (path / ".git").exists():
        commit = git_commit(path)
        return name, {"status": "ok", "commit": commit, "action": "existing", "error": ""}
    if path.exists():
        return name, {"status": "failed", "commit": "", "action": "blocked", "error": "path exists but is not a git repository"}
    path.parent.mkdir(parents=True, exist_ok=True)
    env = os.environ.copy()
    env["GIT_TERMINAL_PROMPT"] = "0"
    env["GIT_SSH_COMMAND"] = env.get("GIT_SSH_COMMAND", "ssh -o BatchMode=yes -o StrictHostKeyChecking=accept-new -o ConnectTimeout=15")
    command = [
        "git",
        "clone",
        "--depth",
        "1",
        "--single-branch",
        item["url"] + ("" if item["url"].endswith(".git") else ".git"),
        str(path),
    ]
    try:
        result = subprocess.run(command, capture_output=True, text=True, timeout=timeout, env=env)
        if result.returncode != 0:
            shutil.rmtree(path, ignore_errors=True)
            error = (result.stderr or result.stdout or "").strip().splitlines()[-1:]
            return name, {"status": "failed", "commit": "", "action": "clone", "error": " ".join(error)[:500]}
        commit = git_commit(path)
        if not commit:
            return name, {"status": "failed", "commit": "", "action": "clone", "error": "clone succeeded but no commit"}
        return name, {"status": "ok", "commit": commit, "action": "clone", "error": ""}
    except subprocess.TimeoutExpired:
        shutil.rmtree(path, ignore_errors=True)
        return name, {"status": "failed", "commit": "", "action": "clone", "error": f"clone timeout after {timeout}s"}
    except Exception as exc:
        shutil.rmtree(path, ignore_errors=True)
        return name, {"status": "failed", "commit": "", "action": "clone", "error": str(exc)[:500]}


def clone_catalog(args):
    manifest_path = Path(args.manifest)
    manifest = read_manifest(manifest_path)
    selected = select_projects(manifest, args.batch)
    todo = [
        (name, item)
        for name, item in selected.items()
        if item.get("status") != "ok" or not (ROOT / item["path"]).exists()
    ]
    if not todo:
        print(json.dumps({"total": 0, "ok": 0, "failed": 0}, ensure_ascii=False))
        return
    print(f"cloning {len(todo)} projects with parallel={args.parallel} timeout={args.timeout}s", flush=True)
    status_path = manifest_path.parent / "clone-status.jsonl"
    status_path.parent.mkdir(parents=True, exist_ok=True)
    with status_path.open("w") as status_file, concurrent.futures.ThreadPoolExecutor(max_workers=args.parallel) as pool:
        futures = {pool.submit(clone_one, name, item, args.timeout): name for name, item in todo}
        for future in concurrent.futures.as_completed(futures):
            name, result = future.result()
            item = manifest["projects"][name]
            item.update(result)
            status_file.write(json.dumps({"name": name, **result}, ensure_ascii=False) + "\n")
            status_file.flush()
            print(f"[{item['status']:6}] {name} {item['commit'][:12]} {result['action']} {result['error']}", flush=True)
    manifest["generated_at"] = now()
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    ok = sum(1 for item in selected.values() if item.get("status") == "ok")
    failed = sum(1 for item in selected.values() if item.get("status") == "failed")
    print(json.dumps({"total": len(selected), "ok": ok, "failed": failed}, ensure_ascii=False))


def plan_catalog(args):
    manifest = read_manifest(args.manifest)
    selected = select_projects(manifest, args.batch)
    ok = sum(1 for item in selected.values() if item.get("status") == "ok")
    pending = sum(1 for item in selected.values() if item.get("status") == "pending")
    failed = sum(1 for item in selected.values() if item.get("status") == "failed")
    print(json.dumps({"total": len(selected), "ok": ok, "pending": pending, "failed": failed}, ensure_ascii=False))
    for name, item in selected.items():
        print(f"{name:36} {item.get('status',''):8} {item.get('commit','')[:12]:12} {item.get('path','')}")


def main():
    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="command", required=True)

    p_import = sub.add_parser("import")
    p_import.add_argument("--xlsx", default=str(DEFAULT_XLSX))
    p_import.add_argument("--out", default=str(DEFAULT_MANIFEST))
    p_import.add_argument("--batch", default="all")
    p_import.set_defaults(func=import_catalog)

    p_clone = sub.add_parser("clone")
    p_clone.add_argument("--manifest", default=str(DEFAULT_MANIFEST))
    p_clone.add_argument("--batch", default="首批候选")
    p_clone.add_argument("--parallel", type=int, default=4)
    p_clone.add_argument("--timeout", type=int, default=600)
    p_clone.set_defaults(func=clone_catalog)

    p_plan = sub.add_parser("plan")
    p_plan.add_argument("--manifest", default=str(DEFAULT_MANIFEST))
    p_plan.add_argument("--batch", default="首批候选")
    p_plan.set_defaults(func=plan_catalog)

    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
